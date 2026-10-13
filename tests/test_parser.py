"""Тесты парсера: MD и DOCX."""
from pathlib import Path

import pytest

from src.ingest.parser import ParseError, parse_file


def test_parse_md(tmp_path: Path):
    f = tmp_path / "doc.md"
    f.write_text("# Заголовок\r\n\r\nТекст раздела.\r\n\r\n\r\nКонец.", encoding="utf-8")
    doc = parse_file(f)
    assert doc.source_type == "md"
    assert "# Заголовок" in doc.text
    assert "\r" not in doc.text          # нормализация переводов строк
    assert doc.page_for_offset(0) == 1   # страница 1 для не-PDF


def test_parse_docx_headings_and_tables(tmp_path: Path):
    from docx import Document

    f = tmp_path / "doc.docx"
    d = Document()
    d.add_heading("Раздел", level=1)
    d.add_paragraph("Обычный абзац.")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Ячейка"
    table.rows[0].cells[1].text = "Значение"
    d.save(f)

    doc = parse_file(f)
    assert doc.source_type == "docx"
    assert "## Раздел" in doc.text       # Heading 1 -> ##
    assert "Обычный абзац." in doc.text
    assert "| Ячейка | Значение |" in doc.text  # таблица -> текст


def test_unsupported_format(tmp_path: Path):
    f = tmp_path / "x.txt"
    f.write_text("text", encoding="utf-8")
    with pytest.raises(ParseError):
        parse_file(f)


def test_scanned_pdf_rejected(tmp_path: Path):
    import fitz

    f = tmp_path / "scan.pdf"
    doc = fitz.open()
    doc.new_page()  # страница без текстового слоя
    doc.save(f)
    doc.close()

    with pytest.raises(ParseError, match="скан"):
        parse_file(f)


def test_text_pdf_page_numbers(tmp_path: Path):
    import fitz

    f = tmp_path / "text.pdf"
    doc = fitz.open()
    for i in range(2):
        page = doc.new_page()
        page.insert_text((72, 72), f"Stranica {i + 1}: tekst dokumenta " * 5)
    doc.save(f)
    doc.close()

    parsed = parse_file(f)
    assert parsed.source_type == "pdf"
    assert parsed.page_for_offset(0) == 1
    # смещение в начале второй страницы -> страница 2
    assert parsed.page_for_offset(parsed.page_starts[-1]) == 2
