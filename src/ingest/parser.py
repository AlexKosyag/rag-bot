"""Парсинг документов: PDF / DOCX / MD -> единый текст + позиции страниц."""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path


class ParseError(Exception):
    """Документ не удалось извлечь (например, PDF-скан без текстового слоя)."""


@dataclass
class ParsedDocument:
    name: str
    source_type: str            # pdf | docx | md
    text: str
    page_starts: list[int] = field(default_factory=lambda: [0])  # смещения страниц
    error: str | None = None

    def page_for_offset(self, offset: int) -> int:
        """1-based номер страницы по смещению символа в text."""
        import bisect

        return bisect.bisect_right(self.page_starts, offset)


_WS_RE = re.compile(r"[ \t]+\n")
_MULTI_NL_RE = re.compile(r"\n{3,}")


def _normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _WS_RE.sub("\n", text)
    return _MULTI_NL_RE.sub("\n\n", text).strip()


def parse_file(path: Path | str) -> ParsedDocument:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".md":
        return _parse_md(path)
    if suffix == ".docx":
        return _parse_docx(path)
    if suffix == ".pdf":
        return _parse_pdf(path)
    raise ParseError(f"Неподдерживаемый формат: {suffix} (нужен PDF/DOCX/MD)")


# --------------------------------------------------------------------- MD
def _parse_md(path: Path) -> ParsedDocument:
    text = _normalize(path.read_text(encoding="utf-8", errors="replace"))
    return ParsedDocument(name=path.name, source_type="md", text=text)


# -------------------------------------------------------------------- PDF
def _parse_pdf(path: Path) -> ParsedDocument:
    import fitz  # pymupdf

    try:
        doc = fitz.open(path)
    except Exception as exc:
        raise ParseError(f"Не удалось открыть PDF: {exc}") from exc

    pages: list[str] = []
    with doc:
        for page in doc:
            pages.append(page.get_text() or "")

    total = sum(len(p.strip()) for p in pages)
    if total < 20 * max(1, len(pages)):
        raise ParseError(
            "PDF без текстового слоя (похоже, скан). OCR вне объёма MVP."
        )

    # Склейка страниц; границы запоминаем для ссылок «стр. N».
    text_parts: list[str] = []
    page_starts: list[int] = []
    offset = 0
    for i, page_text in enumerate(pages):
        page_starts.append(offset)
        chunk = _normalize(page_text)
        text_parts.append(chunk)
        offset += len(chunk) + 2  # + "\n\n"
    text = "\n\n".join(text_parts)
    return ParsedDocument(
        name=path.name, source_type="pdf", text=_normalize(text),
        page_starts=page_starts,
    )


# ------------------------------------------------------------------- DOCX
_HEADING_STYLE_RE = re.compile(r"^(heading|заголовок)\s*(\d)?", re.IGNORECASE)


def _parse_docx(path: Path) -> ParsedDocument:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = Document(path)
    except Exception as exc:
        raise ParseError(f"Не удалось открыть DOCX: {exc}") from exc

    lines: list[str] = []
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            para = Paragraph(child, document)
            style = (para.style.name or "") if para.style is not None else ""
            text = para.text.strip()
            if not text:
                continue
            m = _HEADING_STYLE_RE.match(style)
            if m:
                level = int(m.group(2) or 1)
                lines.append(f"{'#' * min(level + 1, 6)} {text}")
            else:
                lines.append(text)
        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            for row in table.rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                if any(cells):
                    lines.append("| " + " | ".join(cells) + " |")

    text = _normalize("\n".join(lines))
    if not text:
        raise ParseError("DOCX не содержит текста")
    return ParsedDocument(name=path.name, source_type="docx", text=text)
